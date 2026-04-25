"""DOCX report generation service using python-docx."""

import io
import os
import logging
from datetime import datetime, timezone
from typing import List, Dict, Optional, Any

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.services.llm_service import get_llm_provider
from app.services.prompts import (
    SYSTEM_PROMPT,
    PROMPT_INTRODUCAO,
    PROMPT_METODOLOGIA,
    PROMPT_ACHADOS,
    PROMPT_CONCLUSAO,
    PROMPT_RECOMENDACOES,
    PROMPT_LEGENDA,
)

logger = logging.getLogger(__name__)


def _safe(value: Any, default: str = "NÃO INFORMADO") -> str:
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return default
    return str(value)


def _list_to_str(lst: Optional[list], default: str = "NÃO INFORMADO") -> str:
    if not lst:
        return default
    return ", ".join(str(item) for item in lst)


def _add_heading(doc: Document, text: str, level: int = 1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)


def _add_paragraph(doc: Document, text: str, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6)


def _generate_llm_text(prompt_template: str, **kwargs) -> str:
    llm = get_llm_provider()
    prompt = prompt_template.format(**{k: _safe(v) for k, v in kwargs.items()})
    try:
        return llm.generate_text(prompt, system_prompt=SYSTEM_PROMPT)
    except Exception as e:
        logger.error(f"LLM generation failed: {e}")
        return "Texto não disponível — falha na geração automática. Preencher manualmente."


def classify_evidences(evidences: List[Dict]) -> Dict[str, Dict[str, List[Dict]]]:
    """Classify evidences by area and before/after."""
    classified: Dict[str, Dict[str, List[Dict]]] = {}

    for ev in evidences:
        tags = ev.get("tags") or []
        area = "OUTROS"
        phase = "NAO_CLASSIFICADO"

        areas = ["CASCO", "HELICE", "LEME", "TOMADA_DAGUA", "TUBULACAO"]
        for a in areas:
            if a in tags:
                area = a
                break
        if area == "OUTROS":
            filename = (ev.get("filename_original") or "").upper()
            obs = (ev.get("observacao_operador") or "").upper()
            for a in areas:
                if a in filename or a in obs:
                    area = a
                    break

        if "ANTES" in tags:
            phase = "ANTES"
        elif "DEPOIS" in tags:
            phase = "DEPOIS"
        else:
            filename = (ev.get("filename_original") or "").upper()
            obs = (ev.get("observacao_operador") or "").upper()
            if "ANTES" in filename or "ANTES" in obs or "BEFORE" in filename:
                phase = "ANTES"
            elif "DEPOIS" in filename or "DEPOIS" in obs or "AFTER" in filename:
                phase = "DEPOIS"

        if area not in classified:
            classified[area] = {"ANTES": [], "DEPOIS": [], "NAO_CLASSIFICADO": []}
        classified[area][phase].append(ev)

    return classified


def generate_report_docx(
    operation: Dict,
    evidences: List[Dict],
    evidence_images: Dict[str, bytes],
) -> bytes:
    """Generate complete DOCX report."""
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # === CAPA ===
    doc.add_paragraph()
    doc.add_paragraph()
    title = doc.add_heading("RELATÓRIO TÉCNICO DE OPERAÇÃO SUBAQUÁTICA", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    cover_data = [
        ("Cliente", _safe(operation.get("cliente"))),
        ("Navio", _safe(operation.get("navio"))),
        ("IMO", _safe(operation.get("imo"))),
        ("Data da Operação", _safe(operation.get("data_operacao"))),
        ("Local", _safe(operation.get("local"))),
        ("Porto", _safe(operation.get("porto"))),
        ("Tipo de Serviço", _safe(operation.get("tipo_servico"))),
    ]

    table = doc.add_table(rows=len(cover_data), cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(cover_data):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value

    doc.add_page_break()

    # === SUMÁRIO ===
    _add_heading(doc, "SUMÁRIO", level=1)
    sections = [
        "1. Introdução",
        "2. Objetivo do Serviço",
        "3. Escopo da Operação",
        "4. Metodologia Executada",
        "5. Equipe Técnica Envolvida",
        "6. Equipamentos Utilizados",
        "7. Condições Operacionais",
        "8. Evidências Fotográficas",
        "9. Diagnóstico / Achados Técnicos",
        "10. Serviços Executados",
        "11. Conclusão",
        "12. Recomendações Técnicas",
        "13. Anexos",
    ]
    for s in sections:
        _add_paragraph(doc, s)

    doc.add_page_break()

    # === 1. INTRODUÇÃO ===
    _add_heading(doc, "1. INTRODUÇÃO", level=1)
    intro_text = _generate_llm_text(
        PROMPT_INTRODUCAO,
        cliente=operation.get("cliente"),
        navio=operation.get("navio"),
        imo=operation.get("imo"),
        local=operation.get("local"),
        porto=operation.get("porto"),
        data_operacao=operation.get("data_operacao"),
        tipo_servico=operation.get("tipo_servico"),
        descricao_servico=operation.get("descricao_servico"),
    )
    _add_paragraph(doc, intro_text)

    # === 2. OBJETIVO ===
    _add_heading(doc, "2. OBJETIVO DO SERVIÇO", level=1)
    tipo = _safe(operation.get("tipo_servico"))
    desc = _safe(operation.get("descricao_servico"))
    _add_paragraph(
        doc,
        f"O objetivo desta operação foi a realização de {tipo.lower()} "
        f"no navio {_safe(operation.get('navio'))}. {desc}",
    )

    # === 3. ESCOPO ===
    _add_heading(doc, "3. ESCOPO DA OPERAÇÃO", level=1)
    areas = operation.get("areas_inspecionadas") or []
    _add_paragraph(
        doc,
        f"Áreas contempladas nesta operação: {_list_to_str(areas)}.",
    )

    # === 4. METODOLOGIA ===
    _add_heading(doc, "4. METODOLOGIA EXECUTADA", level=1)
    condicoes = operation.get("condicoes") or {}
    met_text = _generate_llm_text(
        PROMPT_METODOLOGIA,
        tipo_servico=operation.get("tipo_servico"),
        areas=_list_to_str(areas),
        equipamentos=_list_to_str(operation.get("equipamentos")),
        visibilidade=condicoes.get("visibilidade"),
        correnteza=condicoes.get("correnteza"),
        mare=condicoes.get("mare"),
        obs_condicoes=condicoes.get("observacoes"),
    )
    _add_paragraph(doc, met_text)

    # === 5. EQUIPE ===
    _add_heading(doc, "5. EQUIPE TÉCNICA ENVOLVIDA", level=1)
    equipe = operation.get("equipe") or []
    if equipe:
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        t.cell(0, 0).text = "Nome"
        t.cell(0, 1).text = "Função"
        t.cell(0, 2).text = "Certificação"
        for member in equipe:
            if isinstance(member, dict):
                row = t.add_row()
                row.cells[0].text = _safe(member.get("nome"))
                row.cells[1].text = _safe(member.get("funcao"))
                row.cells[2].text = _safe(member.get("certificacao"))
    else:
        _add_paragraph(doc, "NÃO INFORMADO")

    # === 6. EQUIPAMENTOS ===
    _add_heading(doc, "6. EQUIPAMENTOS UTILIZADOS", level=1)
    equips = operation.get("equipamentos") or []
    if equips:
        for eq in equips:
            doc.add_paragraph(f"• {eq}", style="List Bullet")
    else:
        _add_paragraph(doc, "NÃO INFORMADO")

    # === 7. CONDIÇÕES OPERACIONAIS ===
    _add_heading(doc, "7. CONDIÇÕES OPERACIONAIS", level=1)
    cond_data = [
        ("Visibilidade", _safe(condicoes.get("visibilidade"))),
        ("Correnteza", _safe(condicoes.get("correnteza"))),
        ("Maré", _safe(condicoes.get("mare"))),
        ("Observações", _safe(condicoes.get("observacoes"))),
    ]
    t = doc.add_table(rows=len(cond_data), cols=2)
    t.style = "Light Grid Accent 1"
    for i, (label, val) in enumerate(cond_data):
        t.cell(i, 0).text = label
        t.cell(i, 1).text = val

    # === 8. EVIDÊNCIAS FOTOGRÁFICAS ===
    _add_heading(doc, "8. EVIDÊNCIAS FOTOGRÁFICAS", level=1)
    classified = classify_evidences(evidences)

    pendencias = []

    if not classified:
        _add_paragraph(doc, "Nenhuma evidência foi associada a esta operação.")
        pendencias.append("Nenhuma evidência fotográfica registrada.")
    else:
        for area, phases in classified.items():
            _add_heading(doc, f"8.{list(classified.keys()).index(area)+1}. {area}", level=2)

            for phase_name in ["ANTES", "DEPOIS", "NAO_CLASSIFICADO"]:
                phase_evs = phases.get(phase_name, [])
                if not phase_evs:
                    continue

                label = {"ANTES": "Antes", "DEPOIS": "Depois", "NAO_CLASSIFICADO": "Não Classificado"}
                _add_heading(doc, f"{label[phase_name]}", level=3)

                for ev in phase_evs:
                    ev_id = ev.get("id", "N/A")
                    ev_key = ev.get("storage_key", "")
                    obs = ev.get("observacao_operador") or ""
                    tags = ev.get("tags") or []

                    if ev_key in evidence_images:
                        try:
                            img_stream = io.BytesIO(evidence_images[ev_key])
                            doc.add_picture(img_stream, width=Inches(5))
                            last_paragraph = doc.paragraphs[-1]
                            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        except Exception as e:
                            logger.warning(f"Could not insert image {ev_key}: {e}")
                            _add_paragraph(doc, f"[Imagem não disponível: {ev.get('filename_original', 'N/A')}]")

                    legenda = _generate_legenda(area, tags, obs)
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(f"Foto ID: {str(ev_id)[:8]} — {legenda}")
                    run.font.size = Pt(9)
                    run.font.italic = True

    # === 9. ACHADOS TÉCNICOS ===
    _add_heading(doc, "9. DIAGNÓSTICO / ACHADOS TÉCNICOS", level=1)
    obs_evidencias = "\n".join(
        f"- {ev.get('observacao_operador', 'Sem observação')}"
        for ev in evidences
        if ev.get("observacao_operador")
    ) or "Nenhuma observação registrada pelos operadores."

    achados_text = _generate_llm_text(
        PROMPT_ACHADOS,
        areas=_list_to_str(areas),
        observacoes_campo=operation.get("observacoes_campo"),
        observacoes_evidencias=obs_evidencias,
    )
    _add_paragraph(doc, achados_text)

    # === 10. SERVIÇOS EXECUTADOS ===
    _add_heading(doc, "10. SERVIÇOS EXECUTADOS", level=1)
    _add_paragraph(doc, _safe(operation.get("descricao_servico")))

    # === 11. CONCLUSÃO ===
    _add_heading(doc, "11. CONCLUSÃO", level=1)
    conclusao_text = _generate_llm_text(
        PROMPT_CONCLUSAO,
        tipo_servico=operation.get("tipo_servico"),
        areas=_list_to_str(areas),
        achados_resumo=achados_text[:500] if achados_text else "Sem achados registrados",
        observacoes_campo=operation.get("observacoes_campo"),
    )
    _add_paragraph(doc, conclusao_text)

    # === 12. RECOMENDAÇÕES ===
    _add_heading(doc, "12. RECOMENDAÇÕES TÉCNICAS", level=1)
    rec_text = _generate_llm_text(
        PROMPT_RECOMENDACOES,
        tipo_servico=operation.get("tipo_servico"),
        achados_resumo=achados_text[:500] if achados_text else "Sem achados registrados",
        recomendacoes_iniciais=operation.get("recomendacoes_iniciais"),
    )
    _add_paragraph(doc, rec_text)

    # === 13. ANEXOS ===
    _add_heading(doc, "13. ANEXOS", level=1)
    _add_heading(doc, "13.1. Tabela de Integridade das Evidências", level=2)

    if evidences:
        t = doc.add_table(rows=1, cols=5)
        t.style = "Light Grid Accent 1"
        headers = ["ID", "Arquivo Original", "SHA256", "Upload", "Área/Tags"]
        for i, h in enumerate(headers):
            t.cell(0, i).text = h

        for ev in evidences:
            row = t.add_row()
            row.cells[0].text = str(ev.get("id", ""))[:8]
            row.cells[1].text = ev.get("filename_original", "N/A")
            row.cells[2].text = ev.get("sha256", "N/A")[:16] + "..."
            row.cells[3].text = str(ev.get("uploaded_at", "N/A"))
            row.cells[4].text = ", ".join(ev.get("tags") or ["N/A"])
    else:
        _add_paragraph(doc, "Nenhuma evidência registrada.")

    # === PENDÊNCIAS ===
    if pendencias:
        doc.add_page_break()
        _add_heading(doc, "PENDÊNCIAS", level=1)
        for p_item in pendencias:
            doc.add_paragraph(f"⚠ {p_item}", style="List Bullet")

    # === RODAPÉ ===
    section = doc.sections[0]
    footer = section.footer
    footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run(
        f"Relatório gerado automaticamente em {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M UTC')} "
        f"— Sistema Agente01"
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(128, 128, 128)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _generate_legenda(area: str, tags: list, observacao: str) -> str:
    try:
        llm = get_llm_provider()
        prompt = PROMPT_LEGENDA.format(
            area=area,
            tags=", ".join(tags) if tags else "N/A",
            observacao=observacao or "Sem observação",
        )
        return llm.generate_text(prompt, system_prompt=SYSTEM_PROMPT)
    except Exception:
        if observacao:
            return observacao
        return "Registro fotográfico da área inspecionada."
